"""
Document parsing and text chunking utilities.
Supports PDF, DOCX, PPTX, CSV, Excel, JSON, TXT, MD, and LOG formats.
"""

import json
import os
from pathlib import Path
from typing import List

import pandas as pd
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from app.core.exceptions import DocumentExtractionError
from app.core.logging_config import get_logger

logger = get_logger(__name__)


def extract_text(file_path: str) -> str:
    """
    Extract readable text content from supported file types.

    Args:
        file_path: Path to the target file.

    Returns:
        Extracted text as a string.

    Raises:
        FileNotFoundError: If the file does not exist.
        DocumentExtractionError: If file format is unsupported or parsing fails.
    """
    path = Path(file_path)
    if not path.is_file():
        logger.error(f"File not found for extraction: {file_path}")
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    logger.info(f"Extracting text from: {path.name} (type: {ext}, size: {path.stat().st_size} bytes)")

    try:
        if ext in [".txt", ".md", ".log"]:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        if ext == ".json":
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                data = json.load(f)
                return json.dumps(data, indent=2)

        if ext == ".csv":
            df = pd.read_csv(path)
            return df.to_string()

        if ext in [".xls", ".xlsx"]:
            df = pd.read_excel(path)
            return df.to_string()

        if ext == ".pdf":
            reader = PdfReader(str(path))
            pages_text = []
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    pages_text.append(page_text)
                else:
                    logger.debug(f"PDF page {idx + 1} produced no text in {path.name}")
            return "\n".join(pages_text)

        if ext == ".docx":
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

        if ext == ".pptx":
            prs = Presentation(str(path))
            slide_texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text)
            return "\n".join(slide_texts)

    except Exception as e:
        logger.error(f"Failed extracting text from {file_path}: {e}", exc_info=True)
        raise DocumentExtractionError(
            f"Failed to extract text from {ext} file: {str(e)}"
        ) from e

    logger.warning(f"Attempted to extract unsupported file format: {ext}")
    raise DocumentExtractionError(
        f"Unsupported file type: {ext}. Supported types are: .txt, .md, .log, .json, .csv, .xls, .xlsx, .pdf, .docx, .pptx"
    )


def chunk_text(text: str, chunk_size: int = 300, overlap: int = 50) -> List[str]:
    """
    Split text into overlapping word-based chunks.

    Args:
        text: Input text string to chunk.
        chunk_size: Maximum number of words per chunk.
        overlap: Number of overlapping words between consecutive chunks.

    Returns:
        List of text chunks.
    """
    if not text or not text.strip():
        return []

    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than 0")
    if overlap < 0:
        raise ValueError("overlap must be non-negative")
    if overlap >= chunk_size:
        raise ValueError(f"overlap ({overlap}) must be strictly less than chunk_size ({chunk_size})")

    words = text.split()
    if not words:
        return []

    chunks: List[str] = []
    start = 0
    step = chunk_size - overlap

    while start < len(words):
        end = start + chunk_size
        chunk = " ".join(words[start:end])
        chunks.append(chunk)
        start += step

    logger.debug(f"Chunked text of {len(words)} words into {len(chunks)} chunks (size: {chunk_size}, overlap: {overlap})")
    return chunks